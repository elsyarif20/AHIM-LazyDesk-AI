import io
import os
import sys
import base64
import csv
import json
import re
from datetime import datetime

from ollama import Client, ResponseError
import keyring
from flask import Flask, jsonify, render_template, request, send_file

app = Flask(__name__)
MAX_UPLOAD_BYTES = 100 * 1024 * 1024  # 100 MB per uploaded file
MAX_PDF_PAGES = 200
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_BYTES
PROVIDERS = {"ollama_cloud", "openai", "gemini"}
MAX_TEXT_CHARS = 60000
KEYRING_SERVICES = {
    "ollama_cloud": ("AHIM LazyDesk AI / Ollama Cloud", "OLLAMA_API_KEY"),
    "openai": ("AHIM LazyDesk AI / OpenAI", "OPENAI_API_KEY"),
    "gemini": ("AHIM LazyDesk AI / Gemini", "GEMINI_API_KEY"),
}

OLLAMA_CLOUD_HOST = "https://ollama.com"
DEFAULT_CLOUD_MODEL = "gemma4:31b"
KEYRING_SERVICE = "AHIM LazyDesk AI / Ollama Cloud"
KEYRING_USERNAME = "OLLAMA_API_KEY"

DEFAULT_PROMPT = """You are AHIM LazyDesk AI, a careful productivity assistant.
Return:
1. Executive summary
2. Action items as checkboxes; do not invent owners or deadlines
3. Important dates and decisions
4. Open questions
Do not invent facts. Mark uncertainty. Reply in the language used by the user."""


def ollama_cloud_call(api_key, model, prompt, content):
    """Call Ollama Cloud using the official Python client; no local Ollama server needed."""
    client = Client(
        host=OLLAMA_CLOUD_HOST,
        headers={"Authorization": "Bearer " + api_key},
    )
    try:
        response = client.chat(
            model=model,
            messages=[
                {"role": "system", "content": prompt},
                {"role": "user", "content": content},
            ],
            stream=False,
        )
    except ResponseError as exc:
        status = getattr(exc, "status_code", None)
        if status in (401, 403):
            raise RuntimeError("Ollama Cloud menolak autentikasi (HTTP 401/403). Periksa OLLAMA_API_KEY.") from None
        if status == 402:
            raise RuntimeError("Ollama Cloud mengembalikan HTTP 402: model ini tidak tersedia pada paket akun. Berhenti dan periksa paket akun Ollama.") from None
        if status == 404:
            raise RuntimeError("Ollama Cloud mengembalikan HTTP 404: model tidak ditemukan. Pastikan model adalah gemma4:31b.") from None
        if status == 429:
            raise RuntimeError("Ollama Cloud mengembalikan HTTP 429: batas penggunaan tercapai. Coba lagi setelah kuota tersedia.") from None
        raise RuntimeError(f"Permintaan Ollama Cloud gagal (HTTP {status or 'unknown'}): {getattr(exc, 'error', str(exc))}") from None

    message = getattr(response, "message", None)
    if message is None and isinstance(response, dict):
        message = response.get("message")
    if isinstance(message, dict):
        content = message.get("content", "")
    else:
        content = getattr(message, "content", "") if message is not None else ""
    result = (content or "").strip()
    if not result:
        raise RuntimeError("Ollama Cloud mengembalikan respons kosong.")
    return result


def provider_key(provider="ollama_cloud"):
    service, username = KEYRING_SERVICES.get(provider, KEYRING_SERVICES["ollama_cloud"])
    env_key = os.getenv(username, "").strip()
    if env_key:
        return env_key
    try:
        return (keyring.get_password(service, username) or "").strip()
    except Exception:
        return ""


def get_api_key():
    return provider_key("ollama_cloud")


def save_provider_key(provider, api_key):
    service, username = KEYRING_SERVICES[provider]
    keyring.set_password(service, username, api_key)


def save_api_key(api_key):
    save_provider_key("ollama_cloud", api_key)


def delete_provider_key(provider):
    service, username = KEYRING_SERVICES[provider]
    try:
        keyring.delete_password(service, username)
    except Exception:
        pass


def delete_saved_api_key():
    delete_provider_key("ollama_cloud")


def provider_chat(provider, api_key, prompt, content, model=None, image_data=None, mime_type=None):
    """Call a configured cloud provider. Image support requires a vision-capable API."""
    import requests
    if provider == "ollama_cloud":
        client = Client(host=OLLAMA_CLOUD_HOST, headers={"Authorization": "Bearer " + api_key})
        message = {"role": "user", "content": content}
        if image_data:
            images = image_data if isinstance(image_data, list) else [image_data]
            message["images"] = [base64.b64decode(img) for img in images]
        try:
            response = client.chat(model=model or DEFAULT_CLOUD_MODEL,
                messages=[{"role": "system", "content": prompt}, message], stream=False)
        except ResponseError as exc:
            raise RuntimeError(f"Ollama Cloud gagal (HTTP {getattr(exc, 'status_code', 'unknown')}). Periksa key, model, dan kuota.") from None
        msg = getattr(response, "message", None)
        result = getattr(msg, "content", "") if msg else ""
        if isinstance(response, dict):
            result = (response.get("message") or {}).get("content", result)
        if not result.strip():
            raise RuntimeError("Provider mengembalikan jawaban kosong.")
        return result.strip(), (model or DEFAULT_CLOUD_MODEL)
    if provider == "openai":
        model = model or os.getenv("OPENAI_MODEL", "gpt-4.1-mini")
        user_content = [{"type": "text", "text": content}]
        if image_data:
            images = image_data if isinstance(image_data, list) else [image_data]
            for img in images:
                user_content.append({"type": "image_url", "image_url": {"url": f"data:{mime_type or 'image/png'};base64,{img}"}})
        response = requests.post("https://api.openai.com/v1/chat/completions",
            headers={"Authorization": "Bearer " + api_key, "Content-Type": "application/json"},
            json={"model": model, "messages": [{"role":"system","content":prompt},{"role":"user","content":user_content}], "temperature":0.2}, timeout=180)
        if not response.ok:
            raise RuntimeError(f"OpenAI API gagal (HTTP {response.status_code}). Periksa API key, model, dan kuota.")
        result = response.json()["choices"][0]["message"]["content"]
        return result.strip(), model
    if provider == "gemini":
        model = model or os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
        parts = [{"text": prompt + "\n\n" + content}]
        if image_data:
            images = image_data if isinstance(image_data, list) else [image_data]
            for img in images:
                parts.append({"inline_data": {"mime_type": mime_type or "image/png", "data": img}})
        response = requests.post(f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}",
            headers={"Content-Type":"application/json"}, json={"contents":[{"parts":parts}], "generationConfig":{"temperature":0.2}}, timeout=180)
        if not response.ok:
            raise RuntimeError(f"Gemini API gagal (HTTP {response.status_code}). Periksa API key, model, dan kuota.")
        candidates = response.json().get("candidates", [])
        result = "".join(p.get("text","") for p in (candidates[0].get("content",{}).get("parts",[]) if candidates else []))
        if not result.strip():
            raise RuntimeError("Gemini tidak mengembalikan teks. Periksa dukungan model dan file.")
        return result.strip(), model
    raise RuntimeError("Provider tidak dikenal.")


def _file_to_text(filename, raw, mimetype=""):
    lower = filename.lower()
    if lower.endswith((".txt", ".md", ".csv", ".json", ".log")):
        return raw.decode("utf-8-sig", errors="replace"), None, None
    if lower.endswith((".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff")):
        allowed = {"image/png":"image/png","image/jpeg":"image/jpeg","image/webp":"image/webp","image/bmp":"image/bmp","image/tiff":"image/tiff"}
        mime = mimetype if mimetype in allowed else ({"jpg":"image/jpeg","jpeg":"image/jpeg","png":"image/png","webp":"image/webp","bmp":"image/bmp","tif":"image/tiff","tiff":"image/tiff"}.get(lower.rsplit(".",1)[-1], "image/png"))
        return "", base64.b64encode(raw).decode("ascii"), mime
    if lower.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            parts.append("\n[TABEL]")
            for row in table.rows:
                parts.append(" | ".join(c.text.replace("\n"," ").strip() for c in row.cells))
        return "\n".join(parts), None, None
    if lower.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"\n[LEMBAR: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                vals = ["" if v is None else str(v) for v in row]
                if any(v.strip() for v in vals):
                    parts.append(" | ".join(vals))
        wb.close()
        return "\n".join(parts), None, None
    if lower.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError("PDF terenkripsi atau dikunci kata sandi.")
        page_count = len(reader.pages)
        if page_count == 0:
            raise ValueError("PDF tidak memiliki halaman.")
        if page_count > MAX_PDF_PAGES:
            raise ValueError(f"PDF berisi {page_count} halaman. Batas maksimum AHIM adalah {MAX_PDF_PAGES} halaman per file.")
        text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
        if text:
            return f"[PDF · {page_count} halaman]\n" + text, None, None
        try:
            import fitz
            pdf = fitz.open(stream=raw, filetype="pdf")
            images = []
            for page in pdf:
                pix = page.get_pixmap(matrix=fitz.Matrix(0.85,0.85), alpha=False)
                images.append(base64.b64encode(pix.tobytes("jpeg")).decode("ascii"))
            pdf.close()
            return f"[PDF SCAN: {page_count} halaman. Baca teks, tabel, dan elemen visual; setiap gambar mewakili satu halaman.]", images, "image/jpeg"
        except Exception as exc:
            raise ValueError(f"PDF berupa scan tetapi gagal dirender. Pastikan PyMuPDF terpasang ({type(exc).__name__}).") from None
    raise ValueError("Format file tidak didukung.")


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/config")
def api_config():
    provider = request.args.get("provider", "ollama_cloud")
    if provider not in PROVIDERS:
        provider = "ollama_cloud"
    key = provider_key(provider)
    return jsonify({"cloudConfigured": bool(key), "model": os.getenv({"ollama_cloud":"OLLAMA_CLOUD_MODEL","openai":"OPENAI_MODEL","gemini":"GEMINI_MODEL"}[provider], {"ollama_cloud":DEFAULT_CLOUD_MODEL,"openai":"gpt-4.1-mini","gemini":"gemini-2.5-flash"}[provider])})


@app.post("/api/settings/key")
def settings_key():
    data = request.get_json(silent=True) or {}
    provider = data.get("provider", "ollama_cloud")
    api_key = (data.get("apiKey") or "").strip()
    if provider not in PROVIDERS:
        return jsonify({"error":"Provider tidak dikenal."}), 400
    if not api_key or len(api_key) > 1000 or "\n" in api_key or "\r" in api_key:
        return jsonify({"error":"API key kosong atau formatnya tidak valid."}), 400
    try:
        reply, model = provider_chat(provider, api_key,
            "You are a connection test. Reply with a short confirmation in Indonesian.",
            "Konfirmasi koneksi AHIM LazyDesk AI.")
        save_provider_key(provider, api_key)
        return jsonify({"ok":True,"model":model,"reply":reply})
    except Exception as exc:
        return jsonify({"error":str(exc) if isinstance(exc, RuntimeError) else f"Tes koneksi gagal ({type(exc).__name__}). Periksa koneksi, model, dan key."}), 502


@app.delete("/api/settings/key")
def remove_key():
    provider = request.args.get("provider", "ollama_cloud")
    if provider not in PROVIDERS:
        return jsonify({"error":"Provider tidak dikenal."}), 400
    delete_provider_key(provider)
    return jsonify({"ok":True,"message":"API key provider dihapus dari penyimpanan lokal."})


@app.post("/api/settings/test")
def test_key():
    data = request.get_json(silent=True) or {}
    provider = data.get("provider", "ollama_cloud")
    if provider not in PROVIDERS:
        return jsonify({"error":"Provider tidak dikenal."}), 400
    key = provider_key(provider)
    if not key:
        return jsonify({"error":"API key belum tersedia."}), 400
    try:
        reply, model = provider_chat(provider, key,
            "You are a connection test. Reply with a short confirmation in Indonesian.",
            "Konfirmasi koneksi AHIM LazyDesk AI.", os.getenv({"ollama_cloud":"OLLAMA_CLOUD_MODEL","openai":"OPENAI_MODEL","gemini":"GEMINI_MODEL"}[provider], None))
        return jsonify({"ok":True,"model":model,"reply":reply})
    except Exception as exc:
        return jsonify({"error":str(exc) if isinstance(exc, RuntimeError) else f"Tes koneksi gagal ({type(exc).__name__}). Periksa koneksi, model, dan key."}), 502


@app.post("/api/summarize")
def summarize():
    data = request.get_json(silent=True) or {}
    text = (data.get("text") or "").strip()
    prompt = (data.get("prompt") or DEFAULT_PROMPT).strip()[:12000]
    provider = data.get("provider", "ollama_cloud")
    upload = request.files.get("file")
    image_data = None
    mime_type = None
    if provider not in PROVIDERS and provider != "demo":
        return jsonify({"error":"Provider tidak dikenal."}), 400
    if upload:
        filename = (upload.filename or "").strip()
        if not filename:
            return jsonify({"error":"Nama file kosong. Pilih file lagi atau tekan Reset Upload."}), 400
        raw = upload.read(MAX_UPLOAD_BYTES + 1)
        if len(raw) > MAX_UPLOAD_BYTES:
            return jsonify({"error":"File melebihi batas 100 MB. Tekan Reset Upload dan pilih file yang lebih kecil."}), 413
        if not raw:
            return jsonify({"error":"File kosong. Pilih file lain atau tekan Reset Upload."}), 400
        try:
            text, image_data, mime_type = _file_to_text(filename, raw, upload.mimetype or "")
        except Exception as exc:
            return jsonify({"error":f"File tidak dapat diproses: {exc} Tekan Reset Upload lalu pilih file lain."}), 400
        if image_data and filename.lower().endswith(".pdf"):
            # Up to 4 rendered pages per cloud vision request; max 200 pages per PDF.
            key = provider_key(provider) if provider in PROVIDERS else ""
            if provider == "demo":
                return jsonify({"error":"Mode demo tidak dapat membaca PDF scan. Pilih provider AI vision online."}), 400
            if not key:
                return jsonify({"error":f"API key {provider} belum diatur. Buka AI Settings untuk menyimpan key.", "needsApiKey":True}), 400
            try:
                page_results = []
                batch_size = 4
                for start_page in range(0, len(image_data), batch_size):
                    batch = image_data[start_page:start_page + batch_size]
                    first_page = start_page + 1
                    last_page = start_page + len(batch)
                    batch_content = (f"{text}\n\nBaca semua gambar halaman PDF terlampir, berurutan dari halaman {first_page} sampai {last_page}. "
                                     "Untuk setiap gambar, tulis judul 'Halaman N', ekstrak teks dan tabel sedekat mungkin dengan sumber, "
                                     "pertahankan urutan kolom, dan tandai bagian yang tidak terbaca. Jangan gabungkan data antarhalaman.")
                    result, model = provider_chat(provider, key, prompt, batch_content, image_data=batch, mime_type=mime_type)
                    page_results.append(result)
                text = "\n\n".join(page_results)
                image_data = None
            except Exception as exc:
                return jsonify({"error":f"Gagal membaca PDF hasil scan: {exc}. Reset Upload lalu coba lagi atau gunakan provider vision lain."}), 502
    if not text and not image_data:
        return jsonify({"error":"Tidak ditemukan teks yang bisa dibaca. Untuk PDF scan atau gambar, pilih provider yang mendukung input visual."}), 400
    if len(text) > MAX_TEXT_CHARS:
        return jsonify({"error":f"Konten terlalu panjang ({len(text):,} karakter). Batas aplikasi 60.000 karakter; pecah dokumen menjadi beberapa bagian."}), 413
    if provider == "demo":
        lines = [line.strip(" -*\t") for line in text.splitlines() if line.strip()]
        result = "## Ringkasan — Mode Demo (bukan hasil AI)\n" + " ".join(lines[:4])[:700] + "\n\n## Tindak lanjut\n- [ ] Tinjau dan verifikasi informasi penting.\n\nMode demo tidak memanggil model AI."
        model = "demo"
    elif provider in PROVIDERS:
        key = provider_key(provider)
        if not key:
            return jsonify({"error":f"API key {provider} belum diatur. Buka AI Settings untuk menyimpan key.", "needsApiKey":True}), 400
        try:
            content = text or "Baca gambar yang dilampirkan. Ekstrak teks, tabel, dan informasi penting secara terstruktur. Tandai teks yang tidak terbaca."
            result, model = provider_chat(provider, key, prompt, content, image_data=image_data, mime_type=mime_type)
        except Exception as exc:
            return jsonify({"error":str(exc) if isinstance(exc, RuntimeError) else f"Gagal memanggil provider ({type(exc).__name__}). Periksa koneksi, key, model, dan kuota."}), 502
    else:
        return jsonify({"error":"Provider tidak dikenal. Pilih Ollama Cloud, OpenAI, Gemini, atau demo."}), 400
    return jsonify({"result":result, "characters":len(text), "backend":provider, "model":model})


def _markdown_rows(content):
    rows = []
    for line in content.splitlines():
        s = line.strip()
        if s.startswith("|") and s.endswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
                continue
            rows.append(cells)
    return rows


@app.post("/api/export")
def export():
    data = request.get_json(silent=True) or {}
    content = (data.get("content") or "").strip()
    fmt = (data.get("format") or "md").lower()
    source_name = (data.get("sourceName") or "").strip()
    if not content:
        return jsonify({"error":"Tidak ada hasil untuk diunduh."}), 400
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    try:
        if fmt == "docx":
            from docx import Document
            from docx.shared import Inches
            doc = Document()
            for line in content.splitlines():
                s = line.strip()
                if not s:
                    continue
                if s.startswith("|") and s.endswith("|"):
                    rows = _markdown_rows(content)
                    if rows:
                        table = doc.add_table(rows=1, cols=len(rows[0]))
                        table.style = "Light Shading Accent 1"
                        for j, value in enumerate(rows[0]): doc.rows[0].cells[j].text = value
                        for row in rows[1:]:
                            cells = table.add_row().cells
                            for j, value in enumerate(row[:len(cells)]): cells[j].text = value
                        break
                if s.startswith("# "): doc.add_heading(s[2:], level=1)
                elif s.startswith("## "): doc.add_heading(s[3:], level=2)
                elif s.startswith("### "): doc.add_heading(s[4:], level=3)
                elif s.startswith(("- ", "* ", "- [ ] ", "- [x] ")): doc.add_paragraph(re.sub(r"^- \[[ xX]\] |^[-*] ", "• ", s), style="List Bullet")
                else: doc.add_paragraph(s.replace("**","").replace("`",""))
            out = io.BytesIO(); doc.save(out); out.seek(0)
            return send_file(out, mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document", as_attachment=True, download_name=f"ahim_lazydesk_{stamp}.docx")
        if fmt == "pdf":
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib import colors
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
            from reportlab.lib.units import mm
            from xml.sax.saxutils import escape
            out = io.BytesIO(); styles = getSampleStyleSheet()
            story = []
            for line in content.splitlines():
                s = line.strip()
                if not s: story.append(Spacer(1, 3*mm)); continue
                if s.startswith("|") and s.endswith("|"):
                    rows = _markdown_rows(content)
                    if rows:
                        t = Table(rows, repeatRows=1, hAlign="LEFT")
                        t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#e8edf4")),("GRID",(0,0),(-1,-1),0.4,colors.grey),("VALIGN",(0,0),(-1,-1),"TOP"),("FONTSIZE",(0,0),(-1,-1),8),("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5)]))
                        story.append(t); story.append(Spacer(1,4*mm)); break
                if s.startswith("# "): story.append(Paragraph(escape(s[2:]), styles["Title"]))
                elif s.startswith("## "): story.append(Paragraph(escape(s[3:]), styles["Heading2"]))
                elif s.startswith("### "): story.append(Paragraph(escape(s[4:]), styles["Heading3"]))
                elif s.startswith(("- ", "* ")): story.append(Paragraph("• "+escape(re.sub(r"^[-*] ","",s).replace("**","")), styles["BodyText"]))
                else: story.append(Paragraph(escape(s.replace("**","").replace("`","")), styles["BodyText"]))
                story.append(Spacer(1, 2*mm))
            SimpleDocTemplate(out, pagesize=A4, rightMargin=18*mm,leftMargin=18*mm,topMargin=18*mm,bottomMargin=18*mm).build(story)
            out.seek(0)
            return send_file(out, mimetype="application/pdf", as_attachment=True, download_name=f"ahim_lazydesk_{stamp}.pdf")
        if fmt == "xlsx":
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
            rows = _markdown_rows(content)
            wb = Workbook(); ws = wb.active; ws.title = "Hasil"
            if rows:
                for row in rows: ws.append(row)
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="17365D")
                ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
                for col in ws.columns:
                    width = min(max(max(len(str(c.value or "")) for c in col)+2, 12), 40)
                    ws.column_dimensions[get_column_letter(col[0].column)].width = width
            else:
                ws.append(["Hasil AI"])
                for line in content.splitlines():
                    if line.strip(): ws.append([line.strip()])
                ws["A1"].font = Font(bold=True, color="FFFFFF"); ws["A1"].fill = PatternFill("solid", fgColor="17365D")
                ws.column_dimensions["A"].width = 90
                for row in ws.iter_rows(): row[0].alignment = Alignment(wrap_text=True, vertical="top")
            out=io.BytesIO(); wb.save(out); out.seek(0)
            return send_file(out, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=f"ahim_lazydesk_{stamp}.xlsx")
        name = f"ahim_lazydesk_{stamp}.md"
        return send_file(io.BytesIO(content.encode("utf-8")), mimetype="text/markdown", as_attachment=True, download_name=name)
    except Exception as exc:
        return jsonify({"error":f"Gagal membuat ekspor {fmt.upper()} ({type(exc).__name__})."}), 500


if __name__ == "__main__":
    if "--test-ollama" in sys.argv:
        key = get_api_key()
        if not key:
            raise SystemExit("API key belum tersedia. Buka AI Settings di aplikasi untuk menambahkan key, lalu jalankan ulang tes.")
        print(ollama_cloud_call(
            key,
            DEFAULT_CLOUD_MODEL,
            "You are a concise assistant. Answer the user's request naturally in one short sentence.",
            "Briefly confirm that you received this message and state what model you are.",
        ))
    else:
        app.run(host=os.getenv("HOST", "127.0.0.1"), port=int(os.getenv("PORT", "5000")), debug=False)
